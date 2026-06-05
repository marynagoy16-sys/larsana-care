"""Extrai texto de PDFs, XLSX e arquivos de texto da pasta plataforma fisio."""
from __future__ import annotations

import hashlib
import re
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

SOURCE = Path(r"c:\projetos\Larsana Care\plataforma fisio")
OUTPUT = Path(r"c:\projetos\Larsana Care\plataforma-fisio-textos-extraidos")

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".tiff"}
TEXT_EXT = {".txt", ".md", ".csv"}
SKIP_EXT = {".py"}


def safe_name(path: Path, root: Path, index: int) -> str:
    rel = path.relative_to(root)
    short = "__".join(rel.parts[-2:]) if len(rel.parts) >= 2 else rel.name
    short = unicodedata.normalize("NFKD", short)
    short = "".join(c if c.isalnum() or c in "._- " else "_" for c in short)
    short = re.sub(r"_+", "_", short).strip("._ ")
    short = short[:80] if short else "arquivo"
    digest = hashlib.md5(str(rel).encode("utf-8")).hexdigest()[:8]
    return f"{index:04d}_{short}_{digest}.txt"


def extract_pdf(path: Path) -> str:
    lines = [f"=== ARQUIVO ORIGINAL: {path} ===", f"=== EXTRAÍDO EM: {datetime.now().isoformat(timespec='seconds')} ===", ""]
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                lines.append("[PDF criptografado — texto não extraído]")
                return "\n".join(lines)

        if not reader.pages:
            lines.append("[PDF sem páginas]")
            return "\n".join(lines)

        for i, page in enumerate(reader.pages, start=1):
            lines.append(f"--- Página {i} ---")
            try:
                text = page.extract_text() or ""
            except Exception as exc:
                text = f"[Erro ao extrair página: {exc}]"
            text = text.strip()
            lines.append(text if text else "[Sem texto extraível nesta página — possível scan/imagem]")
            lines.append("")
    except Exception as exc:
        lines.append(f"[Erro ao ler PDF: {exc}]")
    return "\n".join(lines)


def cell_value(cell) -> str:
    if cell is None:
        return ""
    v = cell.value
    if v is None:
        return ""
    return str(v).strip()


def extract_xlsx(path: Path) -> str:
    lines = [f"=== ARQUIVO ORIGINAL: {path} ===", f"=== EXTRAÍDO EM: {datetime.now().isoformat(timespec='seconds')} ===", ""]
    try:
        wb = load_workbook(str(path), data_only=True, read_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lines.append(f"=== PLANILHA: {sheet_name} ===")
            row_count = 0
            for row in ws.iter_rows():
                values = [cell_value(c) for c in row]
                if any(values):
                    lines.append("\t".join(values))
                    row_count += 1
            if row_count == 0:
                lines.append("[Planilha vazia]")
            lines.append("")
        wb.close()
    except Exception as exc:
        lines.append(f"[Erro ao ler XLSX: {exc}]")
    return "\n".join(lines)


def extract_text_file(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            content = path.read_text(encoding=encoding)
            header = (
                f"=== ARQUIVO ORIGINAL: {path} ===\n"
                f"=== CÓPIA EM: {datetime.now().isoformat(timespec='seconds')} ===\n\n"
            )
            return header + content
        except UnicodeDecodeError:
            continue
    return f"=== ARQUIVO ORIGINAL: {path} ===\n[Erro: encoding não suportado]"


def image_placeholder(path: Path) -> str:
    return (
        f"=== ARQUIVO ORIGINAL: {path} ===\n"
        f"=== TIPO: IMAGEM ({path.suffix.lower()}) ===\n\n"
        "Texto não extraído automaticamente (arquivo de imagem — RG, CNH, foto, carteirinha etc.).\n"
        "Para OCR, instale pytesseract + Tesseract OCR e reexecute com suporte a imagens.\n"
    )


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    index_lines = [
        "ÍNDICE DE EXTRAÇÃO — plataforma fisio",
        f"Gerado em: {datetime.now().isoformat(timespec='seconds')}",
        f"Origem: {SOURCE}",
        f"Destino: {OUTPUT}",
        "",
        "formato\tarquivo_original\tarquivo_extraido\tstatus",
    ]

    stats = {"pdf": 0, "xlsx": 0, "text": 0, "image": 0, "other": 0, "error": 0}

    files = sorted(p for p in SOURCE.rglob("*") if p.is_file())
    file_index = 0
    for src in files:
        ext = src.suffix.lower()
        if ext in SKIP_EXT:
            continue

        file_index += 1
        out_path = OUTPUT / safe_name(src, SOURCE, file_index)
        status = "ok"

        try:
            if ext == ".pdf":
                content = extract_pdf(src)
                stats["pdf"] += 1
            elif ext == ".xlsx":
                content = extract_xlsx(src)
                stats["xlsx"] += 1
            elif ext in TEXT_EXT:
                content = extract_text_file(src)
                stats["text"] += 1
            elif ext in IMAGE_EXT:
                content = image_placeholder(src)
                stats["image"] += 1
            else:
                content = (
                    f"=== ARQUIVO ORIGINAL: {src} ===\n"
                    f"[Tipo {ext} não processado nesta extração]\n"
                )
                stats["other"] += 1

            out_path.write_text(content, encoding="utf-8")
        except Exception as exc:
            status = f"erro: {exc}"
            stats["error"] += 1
            out_path.write_text(
                f"=== ARQUIVO ORIGINAL: {src} ===\n[ERRO NA EXTRAÇÃO: {exc}]\n",
                encoding="utf-8",
            )

        index_lines.append(f"{ext}\t{src}\t{out_path}\t{status}")

    summary = [
        "",
        "=== RESUMO ===",
        f"PDFs extraídos: {stats['pdf']}",
        f"Planilhas XLSX extraídas: {stats['xlsx']}",
        f"Arquivos de texto copiados: {stats['text']}",
        f"Imagens (placeholder): {stats['image']}",
        f"Outros tipos: {stats['other']}",
        f"Erros: {stats['error']}",
        f"Total de arquivos processados: {sum(stats.values())}",
    ]

    index_path = OUTPUT / "_INDICE_EXTRACAO.txt"
    index_path.write_text("\n".join(index_lines + summary), encoding="utf-8")

    print(f"Concluído: {sum(stats.values())} arquivos")
    print(f"Pasta de saída: {OUTPUT}")
    for k, v in stats.items():
        print(f"  {k}: {v}")


if __name__ == "__main__":
    main()
